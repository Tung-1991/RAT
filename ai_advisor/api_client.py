# -*- coding: utf-8 -*-
import json
import os
import shutil
import urllib.error
import urllib.request

from . import history, paths


DEFAULT_PROMPT = """You are an AI Advisor for RAT6. Analyze only the provided internal package.
Read advisor_flow.md first, then advisor_guide inside technical_settings.json before interpreting internal keys.
Do not suggest automatic trading actions. Do not claim web research. Do not tell the bot to edit config.
Act like a trader/risk manager reviewing performance, risk, close reasons, modules, and config history.
When uncertain about an internal key, say what evidence you used instead of inventing behavior."""

TECHNICAL_SETTINGS_LIMIT = 1000000
PROMPT_LIMIT = 200000
ADVISOR_FLOW_LIMIT = 200000
USER_CONTEXT_LIMIT = 100000
PREVIOUS_RESPONSE_LIMIT = 60000
WORKBOOK_LIMIT_ROWS = 80
SUPPORTED_MODELS = ["gpt-5.4-mini", "gpt-5.4", "gpt-5.5"]
MODEL_PRICING_PER_1M = {
    "gpt-5.4-mini": {"input": 0.75, "output": 4.50},
    "gpt-5.4": {"input": 2.50, "output": 15.00},
    "gpt-5.5": {"input": 5.00, "output": 30.00},
}
DEFAULT_MODEL = "gpt-5.4-mini"


DEFAULT_API_SETTINGS = {
    "model": DEFAULT_MODEL,
    "advisor_prompt_limit": PROMPT_LIMIT,
    "advisor_flow_limit": ADVISOR_FLOW_LIMIT,
    "user_context_limit": USER_CONTEXT_LIMIT,
    "technical_settings_limit": TECHNICAL_SETTINGS_LIMIT,
    "previous_response_limit": PREVIOUS_RESPONSE_LIMIT,
    "workbook_limit_rows": WORKBOOK_LIMIT_ROWS,
}


def _safe_int(value, default, min_value=1, max_value=5000000):
    try:
        parsed = int(float(value))
    except Exception:
        return default
    return max(min_value, min(max_value, parsed))


def normalize_model(value):
    model = str(value or DEFAULT_MODEL).strip()
    return model if model in SUPPORTED_MODELS else DEFAULT_MODEL


def _stdout_log(message):
    try:
        print(f"[AI ADVISOR API] {message}", flush=True)
    except Exception:
        pass


def load_api_settings():
    settings = dict(DEFAULT_API_SETTINGS)
    path = paths.advisor_api_settings_path()
    legacy_path = paths.legacy_advisor_api_settings_path()
    source_path = path if os.path.exists(path) else legacy_path
    if os.path.exists(source_path):
        try:
            with open(source_path, "r", encoding="utf-8") as f:
                data = json.load(f)
            if isinstance(data, dict):
                settings.update(data)
        except Exception:
            pass
    settings["technical_settings_limit"] = _safe_int(
        settings.get("technical_settings_limit"),
        TECHNICAL_SETTINGS_LIMIT,
        min_value=1000,
    )
    settings["advisor_prompt_limit"] = _safe_int(
        settings.get("advisor_prompt_limit"),
        PROMPT_LIMIT,
        min_value=1000,
    )
    settings["advisor_flow_limit"] = _safe_int(
        settings.get("advisor_flow_limit"),
        ADVISOR_FLOW_LIMIT,
        min_value=1000,
    )
    settings["user_context_limit"] = _safe_int(
        settings.get("user_context_limit"),
        USER_CONTEXT_LIMIT,
        min_value=1000,
    )
    settings["previous_response_limit"] = _safe_int(
        settings.get("previous_response_limit"),
        PREVIOUS_RESPONSE_LIMIT,
        min_value=0,
    )
    settings["workbook_limit_rows"] = _safe_int(
        settings.get("workbook_limit_rows"),
        WORKBOOK_LIMIT_ROWS,
        min_value=1,
        max_value=10000,
    )
    settings["model"] = normalize_model(settings.get("model"))
    if os.path.exists(legacy_path) and (source_path == legacy_path or os.path.exists(path)):
        try:
            save_api_settings(settings)
            os.remove(legacy_path)
        except Exception:
            pass
    return settings


def save_api_settings(settings):
    paths.ensure_advisor_dirs()
    clean = dict(DEFAULT_API_SETTINGS)
    clean.update(settings or {})
    clean = load_api_settings_from_dict(clean)
    with open(paths.advisor_api_settings_path(), "w", encoding="utf-8") as f:
        json.dump(clean, f, indent=2, ensure_ascii=False)
    return clean


def load_api_settings_from_dict(data):
    data = data or {}
    return {
        "model": normalize_model(data.get("model")),
        "advisor_prompt_limit": _safe_int(
            data.get("advisor_prompt_limit"),
            PROMPT_LIMIT,
            min_value=1000,
        ),
        "advisor_flow_limit": _safe_int(
            data.get("advisor_flow_limit"),
            ADVISOR_FLOW_LIMIT,
            min_value=1000,
        ),
        "user_context_limit": _safe_int(
            data.get("user_context_limit"),
            USER_CONTEXT_LIMIT,
            min_value=1000,
        ),
        "technical_settings_limit": _safe_int(
            data.get("technical_settings_limit"),
            TECHNICAL_SETTINGS_LIMIT,
            min_value=1000,
        ),
        "previous_response_limit": _safe_int(
            data.get("previous_response_limit"),
            PREVIOUS_RESPONSE_LIMIT,
            min_value=0,
        ),
        "workbook_limit_rows": _safe_int(
            data.get("workbook_limit_rows"),
            WORKBOOK_LIMIT_ROWS,
            min_value=1,
            max_value=10000,
        ),
    }


def ensure_advisor_prompt():
    paths.ensure_advisor_dirs()
    path = paths.advisor_prompt_path()
    if not os.path.exists(path):
        template_path = paths.advisor_template_path("advisor_prompt.md")
        if os.path.exists(template_path):
            with open(template_path, "r", encoding="utf-8", errors="replace") as src:
                text = src.read()
        else:
            text = DEFAULT_PROMPT
        with open(path, "w", encoding="utf-8") as f:
            f.write(text)
    return path


def load_advisor_prompt():
    ensure_advisor_prompt()
    text = _read_text(paths.advisor_prompt_path(), limit=load_api_settings()["advisor_prompt_limit"]).strip()
    return text or DEFAULT_PROMPT


def save_advisor_prompt(text):
    paths.ensure_advisor_dirs()
    with open(paths.advisor_prompt_path(), "w", encoding="utf-8") as f:
        f.write((text or DEFAULT_PROMPT).strip() + "\n")
    return paths.advisor_prompt_path()


def _read_text(path, limit=120000):
    if not os.path.exists(path):
        return ""
    with open(path, "r", encoding="utf-8", errors="replace") as f:
        return f.read(limit)


def _workbook_text(limit_rows=None):
    try:
        from openpyxl import load_workbook

        if not os.path.exists(paths.export_path()):
            return ""
        if limit_rows is None:
            limit_rows = load_api_settings().get("workbook_limit_rows", WORKBOOK_LIMIT_ROWS)
        wb = load_workbook(paths.export_path(), data_only=True)
        chunks = []
        for name in wb.sheetnames:
            ws = wb[name]
            chunks.append(f"\n## {name}")
            max_row = min(ws.max_row, limit_rows)
            for row in ws.iter_rows(min_row=1, max_row=max_row, values_only=True):
                chunks.append(" | ".join("" if v is None else str(v) for v in row))
        return "\n".join(chunks)
    except Exception as exc:
        return f"advisor_export.xlsx read warning: {exc}"


def build_api_sections(include_previous_response=False):
    settings = load_api_settings()
    sections = [
        ("advisor_flow.md", _read_text(paths.advisor_flow_path(), limit=settings["advisor_flow_limit"])),
        (
            "technical_settings.json",
            _read_text(paths.technical_settings_path(), limit=settings["technical_settings_limit"]),
        ),
        ("advisor_export.xlsx", _workbook_text(limit_rows=settings["workbook_limit_rows"])),
        ("user_context.md", _read_text(paths.user_context_path(), limit=settings["user_context_limit"])),
    ]
    if include_previous_response:
        sections.append(
            (
                "advisor_response.md",
                _read_text(paths.advisor_response_path(), limit=settings["previous_response_limit"]),
            )
        )
    return sections


def build_api_input(include_previous_response=False):
    sections = []
    for name, text in build_api_sections(include_previous_response=include_previous_response):
        section_name = "previous_advisor_response.md" if name == "advisor_response.md" else name
        sections.extend([f"# {section_name}", text])
    return "\n\n".join(sections)


def estimate_api_payload(include_previous_response=False):
    prompt_text = load_advisor_prompt()
    input_sections = build_api_sections(include_previous_response=include_previous_response)
    model = load_api_settings().get("model", DEFAULT_MODEL)
    pricing = MODEL_PRICING_PER_1M.get(model, MODEL_PRICING_PER_1M[DEFAULT_MODEL])
    text = "\n\n".join(
        part
        for name, section_text in input_sections
        for part in (f"# {'previous_advisor_response.md' if name == 'advisor_response.md' else name}", section_text)
    )
    chars = len(text) + len(prompt_text)
    tokens = max(1, int(chars / 4))
    input_cost = (tokens / 1000000.0) * pricing["input"]
    output_2k_cost = (2000 / 1000000.0) * pricing["output"]
    output_4k_cost = (4000 / 1000000.0) * pricing["output"]
    breakdown = []
    prompt_chars = len(prompt_text)
    breakdown.append(
        {
            "name": "advisor_prompt.md",
            "chars": prompt_chars,
            "tokens": max(1, int(prompt_chars / 4)),
            "included": True,
        }
    )
    for name, section_text in input_sections:
        section_chars = len(section_text or "")
        breakdown.append(
            {
                "name": name,
                "chars": section_chars,
                "tokens": max(1, int(section_chars / 4)),
                "included": True,
            }
        )
    return {
        "chars": chars,
        "tokens": tokens,
        "input_cost_usd": input_cost,
        "estimated_output_2k_usd": output_2k_cost,
        "estimated_output_4k_usd": output_4k_cost,
        "model": model,
        "settings": load_api_settings(),
        "breakdown": breakdown,
    }


def send_package_to_api(prompt=None, include_previous_response=False):
    key = os.environ.get("OPENAI_API_KEY")
    if not key:
        msg = "OPENAI_API_KEY is not configured; API mode skipped."
        _stdout_log(msg)
        history.record_event("advisor_api_missing_key", msg, severity="WARN")
        return {"ok": False, "error": msg}

    settings = load_api_settings()
    model = settings.get("model", DEFAULT_MODEL)
    if model not in SUPPORTED_MODELS:
        msg = f"Unsupported OpenAI model: {model}. Choose one of: {', '.join(SUPPORTED_MODELS)}"
        _stdout_log(msg)
        history.record_event("advisor_api_bad_model", msg, severity="ERROR", payload={"model": model})
        return {"ok": False, "error": msg}
    endpoint = os.environ.get("ADVISOR_API_URL", "https://api.openai.com/v1/responses")
    body_text = build_api_input(include_previous_response=include_previous_response)
    estimate = estimate_api_payload(include_previous_response=include_previous_response)
    _stdout_log(
        "sending "
        f"model={model} endpoint={endpoint} "
        f"chars={estimate.get('chars')} tokens~{estimate.get('tokens')} "
        f"include_response={bool(include_previous_response)}"
    )
    payload = {
        "model": model,
        "instructions": prompt or load_advisor_prompt(),
        "input": body_text,
    }
    req = urllib.request.Request(
        endpoint,
        data=json.dumps(payload).encode("utf-8"),
        headers={
            "Authorization": f"Bearer {key}",
            "Content-Type": "application/json",
        },
        method="POST",
    )
    try:
        with urllib.request.urlopen(req, timeout=90) as resp:
            data = json.loads(resp.read().decode("utf-8"))
        text = data.get("output_text")
        if not text:
            parts = []
            for item in data.get("output", []) or []:
                for content in item.get("content", []) or []:
                    if content.get("type") == "output_text":
                        parts.append(content.get("text", ""))
            text = "\n".join(parts).strip()
        if not text:
            text = json.dumps(data, ensure_ascii=False, indent=2)
        with open(paths.advisor_response_path(), "w", encoding="utf-8") as f:
            f.write(text)
        response_history = paths.advisor_response_history_path()
        os.makedirs(os.path.dirname(response_history), exist_ok=True)
        shutil.copy2(paths.advisor_response_path(), response_history)
        history.record_event(
            "advisor_api_response_saved",
            "Advisor API response saved",
            payload={
                "model": model,
                "response_history": response_history,
                "include_previous_response": bool(include_previous_response),
            },
        )
        _stdout_log(f"response saved response={paths.advisor_response_path()} history={response_history}")
        return {"ok": True, "response": paths.advisor_response_path(), "response_history": response_history, "model": model}
    except urllib.error.HTTPError as exc:
        detail = ""
        try:
            detail = exc.read().decode("utf-8", errors="replace")
        except Exception:
            detail = ""
        msg = f"HTTP {exc.code}: {detail or exc.reason}"
        _stdout_log(msg)
        history.record_event(
            "advisor_api_error",
            msg,
            severity="ERROR",
            payload={"model": model, "endpoint": endpoint, "status": exc.code},
        )
        return {"ok": False, "error": msg}
    except Exception as exc:
        _stdout_log(f"error: {exc}")
        history.record_event("advisor_api_error", str(exc), severity="ERROR", payload={"model": model, "endpoint": endpoint})
        return {"ok": False, "error": str(exc)}
