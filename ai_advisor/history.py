# -*- coding: utf-8 -*-
import csv
import json
import os
import time
from datetime import datetime

from . import config_snapshot, paths


SHEETS = {
    "closed_trades": [
        "Recorded At", "Ticket", "Symbol", "Direction", "Lot", "Entry Time", "Exit Time",
        "Hold Seconds", "Entry Price", "Exit Price", "SL", "TP", "Fee", "Commission",
        "Swap", "Profit", "Close Reason", "Market Mode", "Trigger", "Session ID",
        "Signal Group", "Tactic", "Entry Exit Tactic", "Parent Ticket", "Source Type",
        "MAE", "MFE", "Module Tags", "Config Snapshot ID",
    ],
    "open_trades": [
        "Recorded At", "Ticket", "Symbol", "Direction", "Lot", "Open Time", "Entry Price",
        "SL", "TP", "Profit", "Swap", "Commission", "Tactic", "Entry Exit Tactic",
        "Market Mode", "MAE", "MFE", "Config Snapshot ID",
    ],
    "config_snapshots": [
        "Timestamp", "Snapshot ID", "Reason", "Account ID", "Snapshot JSON",
    ],
    "config_changes": [
        "Timestamp", "Old Snapshot ID", "New Snapshot ID", "Changed Path", "Old Value", "New Value",
    ],
    "events": [
        "Timestamp", "Severity", "Event Type", "Message", "Payload JSON",
    ],
    "summary_daily": ["Key", "Trades", "Profit", "Fee", "Wins", "Losses"],
    "summary_symbol": ["Key", "Trades", "Profit", "Fee", "Wins", "Losses"],
    "summary_timeframe": ["Key", "Trades", "Profit", "Fee", "Wins", "Losses"],
    "summary_signal_group": ["Key", "Trades", "Profit", "Fee", "Wins", "Losses"],
    "summary_close_reason": ["Key", "Trades", "Profit", "Fee", "Wins", "Losses"],
    "summary_module": ["Key", "Trades", "Profit", "Fee", "Wins", "Losses"],
    "trade_config_map": [
        "Recorded At", "Ticket", "Symbol", "Open Time", "Config Snapshot ID", "Source", "Payload JSON",
    ],
}


def _now():
    return datetime.now().isoformat(timespec="seconds")


def _json(value, limit=30000):
    raw = json.dumps(config_snapshot.clone_json(value), ensure_ascii=False, sort_keys=True)
    if len(raw) <= limit:
        return raw
    return raw[:limit] + "...[truncated]"


def _import_openpyxl():
    from openpyxl import Workbook, load_workbook

    return Workbook, load_workbook


def _load_workbook():
    Workbook, load_workbook = _import_openpyxl()
    paths.ensure_advisor_dirs()
    path = paths.history_path()
    if os.path.exists(path):
        wb = load_workbook(path)
    else:
        wb = Workbook()
        default = wb.active
        wb.remove(default)
    for name, headers in SHEETS.items():
        if name not in wb.sheetnames:
            ws = wb.create_sheet(name)
            ws.append(headers)
        else:
            ws = wb[name]
            if ws.max_row == 0:
                ws.append(headers)
    return wb


def _save_workbook(wb):
    path = paths.history_path()
    tmp = f"{path}.tmp.xlsx"
    wb.save(tmp)
    os.replace(tmp, path)


def _append_row(sheet_name, row):
    wb = _load_workbook()
    wb[sheet_name].append(row)
    _save_workbook(wb)


def record_event(event_type, message="", severity="INFO", payload=None):
    try:
        _append_row("events", [_now(), severity, event_type, message, _json(payload or {})])
        return True
    except Exception:
        return False


def _row_values(ws, row_idx):
    return [ws.cell(row_idx, col).value for col in range(1, ws.max_column + 1)]


def _latest_snapshot_from_sheet(ws):
    if ws.max_row < 2:
        return None
    raw = ws.cell(ws.max_row, 5).value
    if not raw or str(raw).endswith("[truncated]"):
        return None
    try:
        return json.loads(raw)
    except Exception:
        return None


def ensure_config_snapshot(reason="observer"):
    try:
        snapshot = config_snapshot.build_snapshot(reason=reason)
        snapshot_id = snapshot.get("config_snapshot_id")
        wb = _load_workbook()
        ws = wb["config_snapshots"]
        existing_ids = {str(ws.cell(r, 2).value) for r in range(2, ws.max_row + 1)}
        latest = _latest_snapshot_from_sheet(ws)
        latest_id = ws.cell(ws.max_row, 2).value if ws.max_row >= 2 else None
        if snapshot_id not in existing_ids:
            ws.append([
                snapshot.get("generated_at"),
                snapshot_id,
                reason,
                snapshot.get("account_id"),
                config_snapshot.compact_snapshot_json(snapshot),
            ])
            if latest_id and latest_id != snapshot_id:
                changes = config_snapshot.diff_snapshots(latest, snapshot) if latest else [
                    ("snapshot_id", latest_id, snapshot_id)
                ]
                ch_ws = wb["config_changes"]
                for path, old_val, new_val in changes:
                    ch_ws.append([
                        _now(),
                        latest_id,
                        snapshot_id,
                        path,
                        _json(old_val, limit=8000),
                        _json(new_val, limit=8000),
                    ])
        _save_workbook(wb)
        return snapshot_id
    except Exception as exc:
        record_event("config_snapshot_error", str(exc), severity="ERROR")
        return "unknown"


def _find_trade_snapshot(ticket):
    try:
        wb = _load_workbook()
        ws = wb["trade_config_map"]
        ticket_str = str(ticket)
        for row in range(ws.max_row, 1, -1):
            if str(ws.cell(row, 2).value) == ticket_str:
                return ws.cell(row, 5).value
    except Exception:
        pass
    return None


def record_trade_opened_data(ticket, symbol="", open_time="", payload=None, source="discovered"):
    try:
        ticket = str(ticket)
        if not ticket:
            return False
        snapshot_id = ensure_config_snapshot(reason=f"trade_open:{ticket}")
        wb = _load_workbook()
        ws = wb["trade_config_map"]
        for row in range(2, ws.max_row + 1):
            if str(ws.cell(row, 2).value) == ticket:
                return False
        if isinstance(open_time, (int, float)):
            open_time = datetime.fromtimestamp(open_time).isoformat(timespec="seconds")
        ws.append([_now(), ticket, symbol, open_time, snapshot_id, source, _json(payload)])
        _save_workbook(wb)
        return True
    except Exception as exc:
        record_event("trade_open_record_error", str(exc), severity="ERROR")
        return False


def record_trade_opened(pos, state=None, market_context=None, source="discovered"):
    ticket = str(getattr(pos, "ticket", ""))
    payload = {
        "magic": getattr(pos, "magic", None),
        "type": getattr(pos, "type", None),
        "volume": getattr(pos, "volume", None),
        "price_open": getattr(pos, "price_open", None),
        "sl": getattr(pos, "sl", None),
        "tp": getattr(pos, "tp", None),
        "state_tactic": (state or {}).get("trade_tactics", {}).get(ticket),
        "market_context": market_context or {},
    }
    return record_trade_opened_data(
        ticket,
        symbol=getattr(pos, "symbol", ""),
        open_time=getattr(pos, "time", ""),
        payload=payload,
        source=source,
    )


def _module_tags(reason, trigger, tactic, session_id):
    text = "|".join(str(v or "").upper() for v in [reason, trigger, tactic, session_id])
    tags = []
    for tag in ["REV_C", "BE_CASH", "TSL", "SL", "GRID", "HEDGE", "DCA", "PCA", "BE", "ANTI_CASH"]:
        if tag in text:
            tags.append(tag)
    return ",".join(tags) if tags else "unknown"


def _source_type(trigger, session_id):
    text = f"{trigger}|{session_id}".upper()
    if "HEDGE" in text:
        return "HEDGE"
    if "GRID" in text:
        return "GRID"
    if "[USER]" in text:
        return "MANUAL"
    if "[BOT]" in text or "AUTO" in text:
        return "BOT"
    return "unknown"


def _parse_signal_group(trigger):
    text = str(trigger or "").upper()
    for group in ["G0", "G1", "G2", "G3"]:
        if group in text:
            return group
    return "unknown"


def _safe_float(value, default=0.0):
    try:
        return float(str(value).replace("$", "").replace(",", "").strip())
    except Exception:
        return default


def record_closed_trade(
    ticket,
    symbol,
    direction,
    volume,
    entry_price,
    sl,
    tp,
    fee,
    pnl,
    close_reason,
    market_mode="ANY",
    trigger_signal="UNK",
    session_id="LEGACY",
    open_time_str="",
    mae_usd=0.0,
    mfe_usd=0.0,
    exit_time=None,
    state=None,
    warn_missing_snapshot=True,
):
    try:
        ticket_str = str(ticket)
        snapshot_id = _find_trade_snapshot(ticket_str)
        if not snapshot_id:
            snapshot_id = ensure_config_snapshot(reason=f"trade_close_missing_open_snapshot:{ticket_str}")
            if warn_missing_snapshot:
                record_event(
                    "missing_trade_open_snapshot",
                    f"No open-time config snapshot for ticket {ticket_str}; using current snapshot.",
                    severity="WARN",
                    payload={"ticket": ticket_str, "snapshot_id": snapshot_id},
                )

        state = state or {}
        tactic = state.get("trade_tactics", {}).get(ticket_str, "unknown")
        ee_tactic = state.get("entry_exit_tactics", {}).get(ticket_str, "unknown")
        parent = state.get("child_to_parent", {}).get(ticket_str, "")
        modules = _module_tags(close_reason, trigger_signal, tactic, session_id)
        exit_time = exit_time or _now()
        hold_seconds = ""
        try:
            if open_time_str:
                hold_seconds = max(0, int((datetime.fromisoformat(exit_time) - datetime.fromisoformat(open_time_str)).total_seconds()))
        except Exception:
            hold_seconds = ""

        row = [
            _now(),
            ticket_str,
            symbol,
            direction,
            volume,
            open_time_str or "",
            exit_time,
            hold_seconds,
            entry_price,
            "",
            sl,
            tp,
            fee,
            fee,
            "",
            pnl,
            close_reason,
            market_mode,
            trigger_signal,
            session_id,
            _parse_signal_group(trigger_signal),
            tactic,
            ee_tactic,
            parent,
            _source_type(trigger_signal, session_id),
            mae_usd,
            mfe_usd,
            modules,
            snapshot_id,
        ]

        wb = _load_workbook()
        ws = wb["closed_trades"]
        replaced = False
        for idx in range(2, ws.max_row + 1):
            if str(ws.cell(idx, 2).value) == ticket_str:
                for col, value in enumerate(row, start=1):
                    ws.cell(idx, col).value = value
                replaced = True
                break
        if not replaced:
            ws.append(row)
        _save_workbook(wb)
        return True
    except Exception as exc:
        record_event("closed_trade_record_error", str(exc), severity="ERROR", payload={"ticket": str(ticket)})
        return False


def sync_from_master_csv():
    try:
        import core.storage_manager as storage_manager

        csv_path = getattr(storage_manager, "MASTER_LOG_FILE", "")
        if not csv_path or not os.path.exists(csv_path):
            record_event("missing_master_trade_csv", "trade_history_master.csv not found", severity="WARN")
            return 0
        count = 0
        with open(csv_path, "r", encoding="utf-8", newline="") as f:
            reader = csv.reader(f)
            header = next(reader, None)
            for row in reader:
                if len(row) < 14:
                    continue
                record_closed_trade(
                    row[1],
                    row[2],
                    row[3],
                    row[4],
                    row[5],
                    row[6],
                    row[7],
                    row[8],
                    row[9],
                    row[10],
                    market_mode=row[11],
                    trigger_signal=row[12],
                    session_id=row[13],
                    open_time_str="",
                    mae_usd=row[14] if len(row) > 14 else 0.0,
                    mfe_usd=row[15] if len(row) > 15 else 0.0,
                    warn_missing_snapshot=False,
                )
                count += 1
        return count
    except Exception as exc:
        record_event("master_csv_sync_error", str(exc), severity="ERROR")
        return 0


def refresh_open_trades(connector=None, state=None, market_contexts=None):
    try:
        wb = _load_workbook()
        ws = wb["open_trades"]
        if ws.max_row > 1:
            ws.delete_rows(2, ws.max_row - 1)
        positions = []
        if connector:
            try:
                positions = connector.get_all_open_positions() or []
            except Exception as exc:
                record_event("open_positions_read_error", str(exc), severity="WARN")
        state = state or {}
        for pos in positions:
            ticket = str(getattr(pos, "ticket", ""))
            snapshot_id = _find_trade_snapshot(ticket) or ensure_config_snapshot(reason=f"open_trade_export:{ticket}")
            direction = "BUY" if int(getattr(pos, "type", 0) or 0) == 0 else "SELL"
            open_time = getattr(pos, "time", "")
            if isinstance(open_time, (int, float)):
                open_time = datetime.fromtimestamp(open_time).isoformat(timespec="seconds")
            ctx = (market_contexts or {}).get(getattr(pos, "symbol", ""), {})
            excursion = state.get("trade_excursions", {}).get(ticket, {})
            ws.append([
                _now(),
                ticket,
                getattr(pos, "symbol", ""),
                direction,
                getattr(pos, "volume", ""),
                open_time,
                getattr(pos, "price_open", ""),
                getattr(pos, "sl", ""),
                getattr(pos, "tp", ""),
                getattr(pos, "profit", ""),
                getattr(pos, "swap", ""),
                getattr(pos, "commission", ""),
                state.get("trade_tactics", {}).get(ticket, "unknown"),
                state.get("entry_exit_tactics", {}).get(ticket, "unknown"),
                ctx.get("market_mode", "unknown") if isinstance(ctx, dict) else "unknown",
                excursion.get("mae_usd", ""),
                excursion.get("mfe_usd", ""),
                snapshot_id,
            ])
        _save_workbook(wb)
        return len(positions)
    except Exception as exc:
        record_event("open_trades_refresh_error", str(exc), severity="ERROR")
        return 0


def rebuild_summaries():
    try:
        wb = _load_workbook()
        closed = wb["closed_trades"]
        headers = [closed.cell(1, col).value for col in range(1, closed.max_column + 1)]
        idx = {name: pos + 1 for pos, name in enumerate(headers)}

        groups = {
            "summary_daily": {},
            "summary_symbol": {},
            "summary_timeframe": {},
            "summary_signal_group": {},
            "summary_close_reason": {},
            "summary_module": {},
        }

        def add(sheet, key, profit, fee):
            item = groups[sheet].setdefault(str(key or "unknown"), {"trades": 0, "profit": 0.0, "fee": 0.0, "wins": 0, "losses": 0})
            item["trades"] += 1
            item["profit"] += profit
            item["fee"] += fee
            if profit >= 0:
                item["wins"] += 1
            else:
                item["losses"] += 1

        for row in range(2, closed.max_row + 1):
            profit = _safe_float(closed.cell(row, idx.get("Profit", 16)).value)
            fee = _safe_float(closed.cell(row, idx.get("Fee", 13)).value)
            exit_time = str(closed.cell(row, idx.get("Exit Time", 7)).value or "")
            date_key = exit_time[:10] if len(exit_time) >= 10 else "unknown"
            add("summary_daily", date_key, profit, fee)
            add("summary_symbol", closed.cell(row, idx.get("Symbol", 3)).value, profit, fee)
            add("summary_timeframe", closed.cell(row, idx.get("Market Mode", 18)).value, profit, fee)
            add("summary_signal_group", closed.cell(row, idx.get("Signal Group", 21)).value, profit, fee)
            add("summary_close_reason", closed.cell(row, idx.get("Close Reason", 17)).value, profit, fee)
            tags = str(closed.cell(row, idx.get("Module Tags", 28)).value or "unknown").split(",")
            for tag in tags:
                add("summary_module", tag.strip() or "unknown", profit, fee)

        for sheet, data in groups.items():
            ws = wb[sheet]
            if ws.max_row > 1:
                ws.delete_rows(2, ws.max_row - 1)
            for key, item in sorted(data.items()):
                ws.append([key, item["trades"], round(item["profit"], 2), round(item["fee"], 2), item["wins"], item["losses"]])
        _save_workbook(wb)
        return True
    except Exception as exc:
        record_event("summary_rebuild_error", str(exc), severity="ERROR")
        return False
