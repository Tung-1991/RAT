# -*- coding: utf-8 -*-
from datetime import datetime, timedelta
from pathlib import Path

from openpyxl import Workbook, load_workbook

from ai_advisor import api_client, history, paths


def _patch_account_dir(monkeypatch, tmp_path):
    monkeypatch.setattr(paths, "account_dir", lambda: str(tmp_path))
    monkeypatch.setattr(paths, "account_id", lambda: "TEST")
    paths.ensure_advisor_dirs()


def _new_history_workbook():
    wb = Workbook()
    wb.remove(wb.active)
    for name, headers in history.SHEETS.items():
        ws = wb.create_sheet(name)
        ws.append(headers)
    return wb


def _row(headers, **values):
    return [values.get(header, "") for header in headers]


def test_advisor_paths_split_history_and_export(monkeypatch, tmp_path):
    _patch_account_dir(monkeypatch, tmp_path)

    assert paths.history_path().replace("\\", "/").endswith("advisor_history.xlsx")
    assert paths.history_path().replace("\\", "/").endswith("history/advisor_history.xlsx")
    assert paths.export_path().replace("\\", "/").endswith("advisor/advisor_export.xlsx")
    assert paths.advisor_response_path().replace("\\", "/").endswith("advisor/advisor_response.md")
    assert paths.advisor_response_history_path().replace("\\", "/").endswith(".md")
    assert "/history/advisor_response_" in paths.advisor_response_history_path().replace("\\", "/")
    assert "/history/user_context_" in paths.user_context_history_path().replace("\\", "/")


def test_storage_csv_paths_live_in_account_history(monkeypatch, tmp_path):
    import core.storage_manager as storage_manager

    repo_root = Path.cwd()
    repo_artifact = repo_root / "data" / "TEST_ACCOUNT"
    if repo_artifact.exists():
        import shutil

        shutil.rmtree(repo_artifact)
    monkeypatch.chdir(tmp_path)
    try:
        storage_manager.set_active_account("TEST_ACCOUNT")

        assert storage_manager.MASTER_LOG_FILE.replace("\\", "/").endswith("data/TEST_ACCOUNT/history/trade_history_master.csv")
        assert storage_manager.HISTORY_FILE.replace("\\", "/").endswith("data/TEST_ACCOUNT/history/trade_history_log.csv")
        assert not repo_artifact.exists()
    finally:
        if repo_artifact.exists():
            import shutil

            shutil.rmtree(repo_artifact)


def test_storage_master_csv_gets_full_time_columns(monkeypatch, tmp_path):
    import csv
    import core.storage_manager as storage_manager

    monkeypatch.chdir(tmp_path)
    history_dir = tmp_path / "data" / "TEST_ACCOUNT" / "history"
    history_dir.mkdir(parents=True)
    master = history_dir / "trade_history_master.csv"
    with open(master, "w", newline="", encoding="utf-8") as f:
        writer = csv.writer(f)
        writer.writerow(["Time", "Ticket", "Symbol", "Type", "Vol", "Entry", "SL", "TP", "Fee", "PnL ($)", "Reason", "Market Mode", "Trigger", "Session_ID", "MAE ($)", "MFE ($)"])
        writer.writerow(["08:34:49 -> 11:12:38", "1", "ETHUSD", "BUY", "1", "1", "1", "1", "0", "10", "Manual_Close", "ANY", "[USER]", "20260609_111238", "0", "10"])

    storage_manager.set_active_account("TEST_ACCOUNT")

    with open(master, "r", newline="", encoding="utf-8") as f:
        rows = list(csv.reader(f))
    header = rows[0]
    row = rows[1]
    assert "Open Time" in header
    assert "Close Time" in header
    assert row[header.index("Open Time")] == "2026-06-09T08:34:49"
    assert row[header.index("Close Time")] == "2026-06-09T11:12:38"


def test_storage_master_csv_keeps_unknown_legacy_close_time_blank(monkeypatch, tmp_path):
    import csv
    import core.storage_manager as storage_manager

    monkeypatch.chdir(tmp_path)
    history_dir = tmp_path / "data" / "TEST_ACCOUNT" / "history"
    history_dir.mkdir(parents=True)
    master = history_dir / "trade_history_master.csv"
    with open(master, "w", newline="", encoding="utf-8") as f:
        writer = csv.writer(f)
        writer.writerow(["Time", "Ticket", "Symbol", "Type", "Vol", "Entry", "SL", "TP", "Fee", "PnL ($)", "Reason", "Market Mode", "Trigger", "Session_ID", "MAE ($)", "MFE ($)", "Open Time", "Close Time"])
        writer.writerow(["13:52:16 -> 16:16:59", "legacy-grid", "ETHUSD", "BUY", "1", "1", "1", "1", "0", "10", "Manual_Close", "ANY", "[GRID]", "GRID", "0", "10", "", "2026-06-10T10:32:39"])

    storage_manager.set_active_account("TEST_ACCOUNT")

    with open(master, "r", newline="", encoding="utf-8") as f:
        rows = list(csv.reader(f))
    header = rows[0]
    row = rows[1]
    assert row[header.index("Open Time")] == ""
    assert row[header.index("Close Time")] == ""


def test_advisor_folder_has_no_nested_dirs_after_export(monkeypatch, tmp_path):
    _patch_account_dir(monkeypatch, tmp_path)
    wb = _new_history_workbook()
    wb.save(paths.history_path())

    result = history.build_export_workbook(export_days=7)

    assert result["ok"] is True
    assert not [p for p in (tmp_path / "advisor").iterdir() if p.is_dir()]


def test_export_workbook_filters_closed_trades_without_touching_full_history(monkeypatch, tmp_path):
    _patch_account_dir(monkeypatch, tmp_path)
    now = datetime.now()
    wb = _new_history_workbook()

    closed_headers = history.SHEETS["closed_trades"]
    wb["closed_trades"].append(
        _row(
            closed_headers,
            **{
                "Recorded At": (now - timedelta(days=3)).isoformat(timespec="seconds"),
                "Ticket": "recent",
                "Symbol": "ETHUSD",
                "Exit Time": (now - timedelta(days=3)).isoformat(timespec="seconds"),
                "Profit": "10",
            },
        )
    )
    wb["closed_trades"].append(
        _row(
            closed_headers,
            **{
                "Recorded At": (now - timedelta(days=20)).isoformat(timespec="seconds"),
                "Ticket": "old",
                "Symbol": "ETHUSD",
                "Exit Time": (now - timedelta(days=20)).isoformat(timespec="seconds"),
                "Profit": "-5",
            },
        )
    )

    open_headers = history.SHEETS["open_trades"]
    wb["open_trades"].append(
        _row(open_headers, **{"Recorded At": now.isoformat(timespec="seconds"), "Ticket": "open-1"})
    )
    wb.save(paths.history_path())

    result_7d = history.build_export_workbook(export_days=7)
    assert result_7d["ok"] is True
    assert result_7d["closed_trades"] == 1
    export_wb = load_workbook(paths.export_path())
    assert [export_wb["closed_trades"].cell(r, 2).value for r in range(2, export_wb["closed_trades"].max_row + 1)] == ["recent"]
    assert export_wb["open_trades"].cell(2, 2).value == "open-1"

    result_30d = history.build_export_workbook(export_days=30)
    assert result_30d["closed_trades"] == 2
    export_wb = load_workbook(paths.export_path())
    assert [export_wb["closed_trades"].cell(r, 2).value for r in range(2, export_wb["closed_trades"].max_row + 1)] == ["recent", "old"]

    full_wb = load_workbook(paths.history_path())
    assert full_wb["closed_trades"].max_row == 3


def test_export_skips_unknown_legacy_closed_trades(monkeypatch, tmp_path):
    _patch_account_dir(monkeypatch, tmp_path)
    now = datetime.now()
    wb = _new_history_workbook()

    closed_headers = history.SHEETS["closed_trades"]
    wb["closed_trades"].append(
        _row(
            closed_headers,
            **{
                "Recorded At": now.isoformat(timespec="seconds"),
                "Ticket": "unknown-date",
                "Symbol": "ETHUSD",
                "Session ID": "GRID",
                "Profit": "10",
            },
        )
    )
    wb["closed_trades"].append(
        _row(
            closed_headers,
            **{
                "Recorded At": now.isoformat(timespec="seconds"),
                "Ticket": "known-date",
                "Symbol": "ETHUSD",
                "Exit Time": now.isoformat(timespec="seconds"),
                "Session ID": "20260610_100000",
                "Profit": "10",
            },
        )
    )
    wb.save(paths.history_path())

    result = history.build_export_workbook(export_days=7)
    assert result["closed_trades"] == 1
    export_wb = load_workbook(paths.export_path())
    assert [export_wb["closed_trades"].cell(r, 2).value for r in range(2, export_wb["closed_trades"].max_row + 1)] == ["known-date"]


def test_master_csv_sync_uses_real_session_date_for_old_rows(monkeypatch, tmp_path):
    import csv
    import core.storage_manager as storage_manager

    monkeypatch.chdir(tmp_path)
    storage_manager.set_active_account("TEST_ACCOUNT")
    with open(storage_manager.MASTER_LOG_FILE, "w", newline="", encoding="utf-8") as f:
        writer = csv.writer(f)
        writer.writerow(["Time", "Ticket", "Symbol", "Type", "Vol", "Entry", "SL", "TP", "Fee", "PnL ($)", "Reason", "Market Mode", "Trigger", "Session_ID", "MAE ($)", "MFE ($)"])
        writer.writerow(["08:34:49 -> 11:12:38", "csv-1", "ETHUSD", "BUY", "1", "1", "1", "1", "0", "10", "Manual_Close", "ANY", "[USER]", "20260609_111238", "0", "10"])

    synced = history.sync_from_master_csv()
    wb = load_workbook(paths.history_path())
    ws = wb["closed_trades"]

    assert synced == 1
    assert ws.cell(2, 6).value == "2026-06-09T08:34:49"
    assert ws.cell(2, 7).value == "2026-06-09T11:12:38"


def test_api_client_reads_advisor_export_not_full_history(monkeypatch, tmp_path):
    _patch_account_dir(monkeypatch, tmp_path)
    full_wb = _new_history_workbook()
    export_wb = _new_history_workbook()
    headers = history.SHEETS["closed_trades"]
    full_wb["closed_trades"].append(_row(headers, **{"Ticket": "full-only"}))
    export_wb["closed_trades"].append(_row(headers, **{"Ticket": "export-only"}))
    full_wb.save(paths.history_path())
    export_wb.save(paths.export_path())

    text = api_client._workbook_text(limit_rows=10)
    assert "export-only" in text
    assert "full-only" not in text


def test_api_client_saves_latest_response_and_history_snapshot(monkeypatch, tmp_path):
    _patch_account_dir(monkeypatch, tmp_path)
    export_wb = _new_history_workbook()
    export_wb.save(paths.export_path())
    paths.user_context_path().replace("\\", "/")
    with open(paths.technical_settings_path(), "w", encoding="utf-8") as f:
        f.write("{}")
    with open(paths.user_context_path(), "w", encoding="utf-8") as f:
        f.write("context")

    class FakeResponse:
        def __enter__(self):
            return self

        def __exit__(self, *_args):
            return False

        def read(self):
            return b'{"output_text":"advisor answer"}'

    monkeypatch.setenv("OPENAI_API_KEY", "test-key")
    monkeypatch.setattr(api_client.urllib.request, "urlopen", lambda *_args, **_kwargs: FakeResponse())

    result = api_client.send_package_to_api()

    assert result["ok"] is True
    assert result["response"].replace("\\", "/").endswith("advisor/advisor_response.md")
    assert result["response_history"].replace("\\", "/").endswith(".md")
    assert "/history/advisor_response_" in result["response_history"].replace("\\", "/")
    assert "advisor_responses" not in result["response_history"].replace("\\", "/")
    with open(paths.advisor_response_path(), "r", encoding="utf-8") as f:
        assert f.read() == "advisor answer"
    with open(result["response_history"], "r", encoding="utf-8") as f:
        assert f.read() == "advisor answer"
